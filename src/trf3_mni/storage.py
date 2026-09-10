import hashlib
import io
import json
import os
import re
import tempfile
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, List, Mapping, Optional

import PyPDF2

from .errors import StorageError
from .models import (
    BatchManifest,
    DocumentReference,
    DownloadedDocument,
    ProcessMetadata,
    ReportWindow,
    StoredArtifact,
)
from src.pgfn_storage import ObjectStorage, StorageLocation


def _safe_filename_part(value: str) -> str:
    cleaned = re.sub(r'[^A-Za-z0-9_-]+', '_', value).strip('_')
    if not cleaned:
        raise StorageError('Metadado inválido para compor o nome do arquivo')
    return cleaned


def _sha256(content: bytes) -> str:
    return hashlib.sha256(content).hexdigest()


def _atomic_write(destination: Path, content: bytes) -> None:
    destination.parent.mkdir(parents=True, exist_ok=True)
    temporary_name = None
    try:
        with tempfile.NamedTemporaryFile(
            mode='wb', prefix=f'.{destination.name}.', suffix='.tmp',
            dir=str(destination.parent), delete=False,
        ) as temporary:
            temporary.write(content)
            temporary.flush()
            os.fsync(temporary.fileno())
            temporary_name = temporary.name
        os.replace(temporary_name, destination)
    except OSError as exc:
        raise StorageError('Falha ao persistir artefato de forma atômica') from exc
    finally:
        if temporary_name and os.path.exists(temporary_name):
            os.unlink(temporary_name)


class PdfStore:
    def __init__(self, output_dir: StorageLocation) -> None:
        self._storage = ObjectStorage(output_dir)

    def save(
        self,
        process: ProcessMetadata,
        document: DownloadedDocument,
        tribunal: str,
        state: str,
        access_link: str,
    ) -> StoredArtifact:
        rendered = self._render_pdf(
            process, document, tribunal=tribunal, state=state, access_link=access_link
        )
        destination = self._destination(process, document.reference, state)
        checksum = _sha256(rendered)

        if self._storage.exists(destination):
            try:
                existing = self._storage.read_bytes(destination)
            except Exception as exc:
                raise StorageError('Falha ao validar PDF já existente') from exc
            if _sha256(existing) == checksum:
                return StoredArtifact(
                    self._storage.location(destination), checksum, len(rendered), reused=True
                )

        location = self._storage.write_bytes(destination, rendered)
        return StoredArtifact(location, checksum, len(rendered), reused=False)

    def find_existing(
        self,
        process: ProcessMetadata,
        reference: DocumentReference,
        tribunal: str,
        state: str,
        access_link: str,
    ) -> Optional[StoredArtifact]:
        """Reutiliza somente um PDF legível cuja identidade esteja comprovada."""
        destination = self._destination(process, reference, state)
        if not self._storage.exists(destination):
            return None
        try:
            content = self._storage.read_bytes(destination)
            reader = PyPDF2.PdfReader(io.BytesIO(content))
            len(reader.pages)  # força a leitura da árvore de páginas
            metadata = reader.metadata or {}
            expected = {
                '/nr_processo': process.process_number,
                '/nr_classe_judicial': process.class_code,
                '/tribunal': tribunal,
                '/secao': state,
                '/dt_protocolo': process.filing_date,
                '/link_acesso': access_link,
                '/id_documento_mni': reference.document_id,
            }
            if any(
                str(metadata.get(key, '')) != str(value)
                for key, value in expected.items()
            ):
                return None
        except Exception:
            return None
        return StoredArtifact(
            self._storage.location(destination),
            _sha256(content),
            len(content),
            reused=True,
        )

    def _destination(
        self,
        process: ProcessMetadata,
        reference: DocumentReference,
        state: str,
    ) -> str:
        filename = '-'.join(
            (
                _safe_filename_part(process.process_number),
                _safe_filename_part(process.class_code),
                _safe_filename_part(state),
                _safe_filename_part(reference.document_id),
            )
        ) + '.pdf'
        return filename

    @staticmethod
    def _render_pdf(
        process: ProcessMetadata,
        document: DownloadedDocument,
        tribunal: str,
        state: str,
        access_link: str,
    ) -> bytes:
        try:
            reader = PyPDF2.PdfReader(io.BytesIO(document.content))
            writer = PyPDF2.PdfWriter()
            for page in reader.pages:
                writer.add_page(page)
            writer.add_metadata({
                '/nr_processo': process.process_number,
                '/nr_classe_judicial': process.class_code,
                '/tribunal': tribunal,
                '/secao': state,
                '/dt_protocolo': process.filing_date,
                '/link_acesso': access_link,
                '/id_documento_mni': document.reference.document_id,
            })
            output = io.BytesIO()
            writer.write(output)
            return output.getvalue()
        except Exception as exc:
            raise StorageError('Conteúdo retornado pelo MNI não é um PDF válido') from exc


class ManifestStore:
    def __init__(self, output_dir: StorageLocation) -> None:
        self._storage = ObjectStorage(output_dir)

    def save(
        self, manifest: BatchManifest, run_id: Optional[str] = None
    ) -> StorageLocation:
        stem = (
            f'manifest-{manifest.window.start.isoformat()}-'
            f'{manifest.window.end.isoformat()}'
        )
        if run_id:
            stem += f'-{_safe_filename_part(run_id)}'
        filename = f'{stem}.json'
        content = json.dumps(
            manifest.to_dict(), ensure_ascii=False, indent=2, sort_keys=True
        ).encode('utf-8')
        return self._storage.write_bytes(filename, content)


class ReportStore:
    """Persiste a resposta original e uma visão XLSX deduplicada por processo."""

    def __init__(self, output_dir: StorageLocation) -> None:
        self._storage = ObjectStorage(output_dir)

    def save(
        self,
        report: Mapping[str, Any],
        window: ReportWindow,
    ) -> List[StoredArtifact]:
        stem = f'relatorio-processos-trf3-{window.start}-{window.end}'
        json_content = json.dumps(
            report, ensure_ascii=False, indent=2, sort_keys=True
        ).encode('utf-8')
        xlsx_content = self._render_xlsx(report)
        return [
            self._save_content(f'{stem}.json', json_content),
            self._save_content(f'{stem}.xlsx', xlsx_content),
        ]

    @staticmethod
    def _render_xlsx(report: Mapping[str, Any]) -> bytes:
        from openpyxl import Workbook

        raw_rows = report.get('Resultado')
        if not isinstance(raw_rows, list):
            raise StorageError('Relatório não possui linhas para gerar a planilha')
        rows = []
        seen_processes = set()
        headers = []
        for raw in raw_rows:
            if not isinstance(raw, dict):
                raise StorageError('Relatório possui linha inválida para a planilha')
            normalized = re.sub(r'\D', '', str(raw.get('nr_processo', '')))
            if not normalized or normalized in seen_processes:
                continue
            seen_processes.add(normalized)
            rows.append(raw)
            for key in raw:
                if key not in headers:
                    headers.append(key)

        workbook = Workbook()
        workbook.properties.created = datetime(2000, 1, 1, tzinfo=timezone.utc)
        workbook.properties.modified = datetime(2000, 1, 1, tzinfo=timezone.utc)
        sheet = workbook.active
        sheet.title = 'Processos TRF3'
        sheet.freeze_panes = 'A2'
        sheet.append(headers)
        if headers:
            sheet.auto_filter.ref = (
                f'A1:{_excel_column(len(headers))}{len(rows) + 1}'
            )
        for row in rows:
            sheet.append([_excel_value(row.get(header)) for header in headers])
        output = io.BytesIO()
        workbook.save(output)
        return output.getvalue()

    def _save_content(self, destination: str, content: bytes) -> StoredArtifact:
        checksum = _sha256(content)
        if self._storage.exists(destination):
            try:
                existing = self._storage.read_bytes(destination)
            except Exception as exc:
                raise StorageError('Falha ao validar relatório já existente') from exc
            if _sha256(existing) == checksum:
                return StoredArtifact(
                    self._storage.location(destination), checksum, len(existing), reused=True
                )
        location = self._storage.write_bytes(destination, content)
        return StoredArtifact(location, checksum, len(content), reused=False)


class FailureReportStore:
    """Exporta uma fila operacional a partir das falhas de um manifesto."""

    ACTIONS = {
        'IntegrationError': (
            'retry_transient',
            'Reexecutar seletivamente; investigar conectividade se persistir.',
        ),
        'InvalidResponseError': (
            'retry_transient',
            'Reexecutar seletivamente; validar contrato da resposta se persistir.',
        ),
        'StorageError': (
            'retry_once',
            'Reexecutar uma vez; investigar conteúdo/armazenamento se persistir.',
        ),
        'ProcessNotFoundError': (
            'source_reconciliation',
            'Conferir divergência entre relatório gerencial e MNI antes de repetir.',
        ),
        'DocumentSelectionError': (
            'business_review',
            'Validar regra funcional e disponibilidade do documento tipo 58.',
        ),
        'UnexpectedError': (
            'technical_review',
            'Investigar logs e causa técnica antes de reexecutar.',
        ),
    }
    DEFAULT_ACTION = (
        'technical_review',
        'Investigar a causa antes de autorizar uma reexecução.',
    )

    def __init__(self, output_dir: StorageLocation) -> None:
        self._storage = ObjectStorage(output_dir)

    def save(
        self,
        manifest: Mapping[str, Any],
        run_id: Optional[str] = None,
    ) -> List[StoredArtifact]:
        window = manifest.get('window')
        results = manifest.get('results')
        if not isinstance(window, dict) or not window.get('start') or not window.get('end'):
            raise StorageError('Manifesto não possui janela para relatório de falhas')
        if not isinstance(results, list):
            raise StorageError('Manifesto não possui resultados para relatório de falhas')
        rows = []
        for result in results:
            if not isinstance(result, dict):
                raise StorageError('Manifesto possui resultado inválido')
            if result.get('status') != 'failed':
                continue
            error_type = str(result.get('error_type') or 'UnknownError')
            action_code, recommendation = self.ACTIONS.get(
                error_type,
                self.DEFAULT_ACTION,
            )
            rows.append({
                'process_number': str(result.get('process_number') or ''),
                'error_type': error_type,
                'error_message': str(result.get('error_message') or ''),
                'action_code': action_code,
                'recommended_action': recommendation,
                'started_at': result.get('started_at'),
                'finished_at': result.get('finished_at'),
                'duration_seconds': result.get('duration_seconds'),
            })
        rows.sort(key=lambda item: (item['error_type'], item['process_number']))
        by_error_type = []
        for error_type in sorted({item['error_type'] for item in rows}):
            matching = [item for item in rows if item['error_type'] == error_type]
            by_error_type.append({
                'error_type': error_type,
                'count': len(matching),
                'action_code': matching[0]['action_code'],
                'recommended_action': matching[0]['recommended_action'],
            })
        payload = {
            'window': {'start': window['start'], 'end': window['end']},
            'source_run_id': manifest.get('run_id'),
            'source_manifest': manifest.get('source_manifest'),
            'summary': {
                'total_failures': len(rows),
                'by_error_type': by_error_type,
            },
            'failures': rows,
        }
        stem = f'pendencias-trf3-{window["start"]}-{window["end"]}'
        if run_id:
            stem += f'-{_safe_filename_part(run_id)}'
        json_content = json.dumps(
            payload, ensure_ascii=False, indent=2, sort_keys=True
        ).encode('utf-8')
        xlsx_content = self._render_xlsx(rows, by_error_type)
        return [
            self._save_content(f'{stem}.json', json_content),
            self._save_content(f'{stem}.xlsx', xlsx_content),
        ]

    @staticmethod
    def _render_xlsx(
        rows: List[Dict[str, Any]],
        summary: List[Dict[str, Any]],
    ) -> bytes:
        from openpyxl import Workbook

        workbook = Workbook()
        workbook.properties.created = datetime(2000, 1, 1, tzinfo=timezone.utc)
        workbook.properties.modified = datetime(2000, 1, 1, tzinfo=timezone.utc)
        pending = workbook.active
        pending.title = 'Pendencias'
        headers = [
            'process_number', 'error_type', 'error_message', 'action_code',
            'recommended_action', 'started_at', 'finished_at', 'duration_seconds',
        ]
        pending.append(headers)
        pending.freeze_panes = 'A2'
        for row in rows:
            pending.append([_excel_value(row.get(header)) for header in headers])
        if headers:
            pending.auto_filter.ref = (
                f'A1:{_excel_column(len(headers))}{len(rows) + 1}'
            )

        summary_sheet = workbook.create_sheet('Resumo')
        summary_headers = [
            'error_type', 'count', 'action_code', 'recommended_action'
        ]
        summary_sheet.append(summary_headers)
        summary_sheet.freeze_panes = 'A2'
        for item in summary:
            summary_sheet.append([
                _excel_value(item.get(header)) for header in summary_headers
            ])
        output = io.BytesIO()
        workbook.save(output)
        return output.getvalue()

    def _save_content(self, destination: str, content: bytes) -> StoredArtifact:
        checksum = _sha256(content)
        if self._storage.exists(destination):
            try:
                existing = self._storage.read_bytes(destination)
            except Exception as exc:
                raise StorageError('Falha ao validar relatório de pendências') from exc
            if _sha256(existing) == checksum:
                return StoredArtifact(
                    self._storage.location(destination), checksum, len(existing), reused=True
                )
        location = self._storage.write_bytes(destination, content)
        return StoredArtifact(location, checksum, len(content), reused=False)


def _excel_value(value: Any) -> Any:
    if value is None or isinstance(value, (bool, int, float)):
        return value
    if not isinstance(value, str):
        value = json.dumps(value, ensure_ascii=False, sort_keys=True)
    if value.startswith(('=', '+', '-', '@')):
        value = "'" + value
    return value[:32767]


def _excel_column(index: int) -> str:
    if index <= 0:
        return 'A'
    result = ''
    while index:
        index, remainder = divmod(index - 1, 26)
        result = chr(65 + remainder) + result
    return result
