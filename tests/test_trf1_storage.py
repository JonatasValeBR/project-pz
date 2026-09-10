import json
import unittest
from datetime import date
from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import load_workbook

from src.trf1_scraping.storage import COLUMNS, save_process_list
from src.pgfn_storage import ObjectStorage


class Trf1StorageTest(unittest.TestCase):
    def test_saves_deterministic_json_and_xlsx(self) -> None:
        records = [
            {
                "id_processo": "2",
                "nr_processo": "0002",
                "nr_classe_judicial": "120",
                "secao": "df",
                "dt_autuacao": "29/07/2026",
                "link_acesso": "https://example.invalid",
                "filtro": "FAZENDA",
            },
            {
                "id_processo": "1",
                "nr_processo": "0001",
                "nr_classe_judicial": "7",
                "secao": "go",
                "dt_autuacao": "29/07/2026",
                "link_acesso": "https://example.invalid",
                "filtro": "Delegado da Receita Federal",
            },
        ]
        with TemporaryDirectory() as temporary:
            summary = save_process_list(records, date(2026, 7, 29), Path(temporary))
            payload = json.loads(Path(summary["json_path"]).read_text("utf-8"))
            self.assertEqual([item["nr_processo"] for item in payload["processos"]],
                             ["0001", "0002"])
            workbook = load_workbook(summary["xlsx_path"], read_only=True)
            rows = list(workbook["Processos"].iter_rows(values_only=True))
            workbook.close()
            self.assertEqual(rows[0], COLUMNS)
            self.assertEqual(len(rows), 3)

    def test_rerun_replaces_same_logical_date(self) -> None:
        with TemporaryDirectory() as temporary:
            root = Path(temporary)
            first = save_process_list([], date(2026, 7, 29), root)
            second = save_process_list([], date(2026, 7, 29), root)
            self.assertEqual(first["json_path"], second["json_path"])
            self.assertEqual(first["xlsx_path"], second["xlsx_path"])

    def test_saves_to_memory_bucket(self) -> None:
        root = "memory://trf1-storage-test/reports"
        summary = save_process_list([], date(2026, 7, 29), root)
        storage = ObjectStorage(root)

        self.assertTrue(str(summary["json_path"]).startswith("memory://"))
        self.assertTrue(
            storage.exists("lista-processos-trf1-2026-07-29.json")
        )
