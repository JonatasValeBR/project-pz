import io
import json
import tempfile
import unittest
from datetime import date, datetime, timezone
from pathlib import Path

import PyPDF2

from src.trf3_mni.models import (
    BatchManifest,
    DocumentReference,
    DownloadedDocument,
    ProcessMetadata,
    ProcessResult,
    ReportWindow,
)
from src.trf3_mni.storage import (
    FailureReportStore,
    ManifestStore,
    PdfStore,
    ReportStore,
)


def blank_pdf() -> bytes:
    writer = PyPDF2.PdfWriter()
    writer.add_blank_page(width=100, height=100)
    output = io.BytesIO()
    writer.write(output)
    return output.getvalue()


class StorageTest(unittest.TestCase):
    def test_failure_report_contains_only_failures_and_actions(self) -> None:
        from openpyxl import load_workbook

        manifest = {
            'window': {'start': '2026-07-15', 'end': '2026-07-15'},
            'results': [
                {
                    'process_number': 'processo-1',
                    'status': 'failed',
                    'error_type': 'StorageError',
                    'error_message': '=mensagem',
                },
                {'process_number': 'processo-2', 'status': 'success'},
            ],
        }
        with tempfile.TemporaryDirectory() as temp_dir:
            artifacts = FailureReportStore(Path(temp_dir)).save(manifest)
            json_path = next(
                item.path for item in artifacts if item.path.suffix == '.json'
            )
            xlsx_path = next(
                item.path for item in artifacts if item.path.suffix == '.xlsx'
            )
            payload = json.loads(json_path.read_text())
            workbook = load_workbook(xlsx_path, read_only=True)
            rows = list(workbook['Pendencias'].iter_rows(values_only=True))

        self.assertEqual(payload['summary']['total_failures'], 1)
        self.assertEqual(payload['failures'][0]['action_code'], 'retry_once')
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[1][2], "'=mensagem")

    def test_report_store_writes_json_and_deduplicated_xlsx(self) -> None:
        from openpyxl import load_workbook

        report = {
            'Resultado': [
                {'nr_processo': '0001-00', 'descricao': '=formula'},
                {'nr_processo': '0001.00', 'descricao': 'duplicado'},
                {'nr_processo': '0002-00', 'descricao': 'segundo'},
            ]
        }
        window = ReportWindow(date(2026, 7, 15), date(2026, 7, 15))
        with tempfile.TemporaryDirectory() as temp_dir:
            artifacts = ReportStore(Path(temp_dir)).save(report, window)
            workbook = load_workbook(
                next(item.path for item in artifacts if item.path.suffix == '.xlsx'),
                read_only=True,
            )
            rows = list(workbook.active.iter_rows(values_only=True))

        self.assertEqual(len(artifacts), 2)
        self.assertEqual(len(rows), 3)
        self.assertEqual(rows[1][1], "'=formula")

    def test_pdf_write_is_idempotent_by_checksum(self) -> None:
        process = ProcessMetadata(
            'processo-sintetico',
            '2026-07-15',
            '100',
            '1',
            [DocumentReference('documento-sintetico', '58')],
        )
        document = DownloadedDocument(process.documents[0], blank_pdf())

        with tempfile.TemporaryDirectory() as temp_dir:
            store = PdfStore(Path(temp_dir))
            first = store.save(process, document, 'TRF 3', 'sp', 'https://example.test')
            second = store.save(process, document, 'TRF 3', 'sp', 'https://example.test')

            self.assertFalse(first.reused)
            self.assertTrue(second.reused)
            self.assertEqual(first.sha256, second.sha256)
            self.assertEqual(list(Path(temp_dir).glob('*.tmp')), [])

            existing = store.find_existing(
                process,
                process.documents[0],
                'TRF 3',
                'sp',
                'https://example.test',
            )
            self.assertIsNotNone(existing)
            self.assertTrue(existing.reused)
            self.assertEqual(existing.sha256, first.sha256)

            mismatched = store.find_existing(
                process,
                process.documents[0],
                'TRF 3',
                'sp',
                'https://outro-link.example.test',
            )
            self.assertIsNone(mismatched)

            first.path.write_bytes(b'pdf-corrompido')
            corrupted = store.find_existing(
                process,
                process.documents[0],
                'TRF 3',
                'sp',
                'https://example.test',
            )
            self.assertIsNone(corrupted)

    def test_pdf_store_supports_memory_bucket(self) -> None:
        process = ProcessMetadata(
            'processo-memory',
            '2026-07-15',
            '100',
            '1',
            [DocumentReference('documento-memory', '58')],
        )
        document = DownloadedDocument(process.documents[0], blank_pdf())
        store = PdfStore('memory://trf3-storage-test/pdfs')

        first = store.save(process, document, 'TRF 3', 'sp', 'https://example.test')
        second = store.save(process, document, 'TRF 3', 'sp', 'https://example.test')

        self.assertTrue(str(first.path).startswith('memory://'))
        self.assertFalse(first.reused)
        self.assertTrue(second.reused)

    def test_manifest_contains_summary(self) -> None:
        instant = datetime(2026, 7, 16, tzinfo=timezone.utc)
        manifest = BatchManifest(
            ReportWindow(date(2026, 7, 15), date(2026, 7, 15)),
            instant,
            instant,
            [ProcessResult('processo-sintetico', 'failed', error_type='Known')],
        )
        with tempfile.TemporaryDirectory() as temp_dir:
            path = ManifestStore(Path(temp_dir)).save(manifest)
            content = json.loads(path.read_text())
        self.assertEqual(content['summary']['failed'], 1)
        self.assertIn('duration_seconds', content['results'][0])

    def test_versioned_manifest_preserves_base_manifest(self) -> None:
        instant = datetime(2026, 7, 16, tzinfo=timezone.utc)
        manifest = BatchManifest(
            ReportWindow(date(2026, 7, 15), date(2026, 7, 15)),
            instant,
            instant,
            [],
        )
        with tempfile.TemporaryDirectory() as temp_dir:
            store = ManifestStore(Path(temp_dir))
            original = store.save(manifest)
            recovery = store.save(manifest, run_id='retry/manual:01')

            self.assertTrue(original.is_file())
            self.assertTrue(recovery.is_file())
            self.assertNotEqual(original, recovery)
